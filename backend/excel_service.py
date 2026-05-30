import os
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from deepseek_text import generate_product_copy


DEFAULT_MAX_EXCEL_ROWS = 20
INPUT_HEADER_ALIASES = {
    "product_name": ("商品名称", "product_name"),
    "product_info": ("商品信息", "product_info"),
    "target_platform": ("目标平台", "target_platform"),
    "tone": ("文案语气", "tone"),
    "language": ("输出语言", "language"),
}
OUTPUT_HEADERS = [
    "商品名称",
    "商品信息",
    "目标平台",
    "文案语气",
    "输出语言",
    "生成标题",
    "生成卖点",
    "生成详情页文案",
    "生成搜索关键词",
    "生成短视频口播文案",
    "处理状态",
    "错误信息",
]
COLUMN_WIDTHS = {
    "商品名称": 24,
    "商品信息": 36,
    "目标平台": 14,
    "文案语气": 18,
    "输出语言": 12,
    "生成标题": 36,
    "生成卖点": 42,
    "生成详情页文案": 60,
    "生成搜索关键词": 36,
    "生成短视频口播文案": 48,
    "处理状态": 12,
    "错误信息": 36,
}


class ExcelValidationError(ValueError):
    pass


def generate_excel_file(file_stream) -> BytesIO:
    workbook = load_workbook(file_stream, data_only=True)
    sheet = workbook.worksheets[0]
    header_map = _build_header_map(sheet)

    if "product_name" not in header_map:
        raise ExcelValidationError("Excel 必须包含“商品名称”或“product_name”列")

    output_workbook = Workbook()
    output_sheet = output_workbook.active
    output_sheet.title = "商品文案结果"
    output_sheet.append(OUTPUT_HEADERS)

    processed_count = 0
    max_rows = _get_max_excel_rows()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        product_name = _get_row_value(row, header_map, "product_name")

        if not product_name:
            continue

        if processed_count >= max_rows:
            break

        row_data = {
            "product_name": product_name,
            "product_info": _get_row_value(row, header_map, "product_info"),
            "target_platform": _get_row_value(row, header_map, "target_platform") or "淘宝",
            "tone": _get_row_value(row, header_map, "tone") or "简洁、有购买欲",
            "language": _get_row_value(row, header_map, "language") or "中文",
        }

        output_sheet.append(_build_output_row(row_data))
        processed_count += 1

    _apply_output_layout(output_sheet)

    output = BytesIO()
    output_workbook.save(output)
    output.seek(0)
    return output


def _build_header_map(sheet) -> dict:
    header_map = {}
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])

    for index, header in enumerate(first_row):
        header_name = _cell_to_text(header)
        column_name = _normalize_input_header(header_name)
        if column_name and column_name not in header_map:
            header_map[column_name] = index

    return header_map


def _get_row_value(row, header_map: dict, column_name: str) -> str:
    column_index = header_map.get(column_name)

    if column_index is None or column_index >= len(row):
        return ""

    return _cell_to_text(row[column_index])


def _build_output_row(row_data: dict) -> list:
    generated = {
        "title": "",
        "selling_points": "",
        "description": "",
        "keywords": "",
        "short_video_script": "",
    }
    status = "成功"
    error_message = ""

    try:
        result = generate_product_copy(
            product_name=row_data["product_name"],
            product_info=row_data["product_info"],
            target_platform=row_data["target_platform"],
            tone=row_data["tone"],
            language=row_data["language"],
        )
        generated["title"] = _cell_to_text(result.get("title"))
        generated["selling_points"] = _join_list_value(result.get("selling_points"))
        generated["description"] = _cell_to_text(result.get("description"))
        generated["keywords"] = _join_list_value(result.get("keywords"))
        generated["short_video_script"] = _cell_to_text(result.get("short_video_script"))
    except Exception as error:
        status = "失败"
        error_message = str(error)

    return [
        row_data["product_name"],
        row_data["product_info"],
        row_data["target_platform"],
        row_data["tone"],
        row_data["language"],
        generated["title"],
        generated["selling_points"],
        generated["description"],
        generated["keywords"],
        generated["short_video_script"],
        status,
        error_message,
    ]


def _join_list_value(value) -> str:
    if value is None:
        return ""

    if isinstance(value, list):
        return "；".join(_cell_to_text(item) for item in value if _cell_to_text(item))

    return _cell_to_text(value)


def _normalize_input_header(header_name: str) -> str:
    for column_name, aliases in INPUT_HEADER_ALIASES.items():
        if header_name in aliases:
            return column_name

    return ""


def _apply_output_layout(sheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="F3F4F6")
    header_font = Font(bold=True)
    top_wrap_alignment = Alignment(vertical="top", wrap_text=True)

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = top_wrap_alignment

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = top_wrap_alignment

    for index, header in enumerate(OUTPUT_HEADERS, start=1):
        column_letter = get_column_letter(index)
        sheet.column_dimensions[column_letter].width = COLUMN_WIDTHS[header]

    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(OUTPUT_HEADERS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"


def _cell_to_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def _get_max_excel_rows() -> int:
    value = os.getenv("MAX_EXCEL_ROWS", str(DEFAULT_MAX_EXCEL_ROWS))

    try:
        max_rows = int(value)
    except ValueError:
        return DEFAULT_MAX_EXCEL_ROWS

    if max_rows < 1:
        return DEFAULT_MAX_EXCEL_ROWS

    return max_rows
